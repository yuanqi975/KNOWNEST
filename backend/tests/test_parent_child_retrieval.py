from __future__ import annotations

import asyncio
from pathlib import Path
from types import SimpleNamespace

import pytest
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.db import Base
from app.models.models import Document, DocumentChunk, KnowledgeBase, User, Workspace
from app.services import ingestion, retrieval
from app.services.parser import ChunkingConfig, extract_text


@pytest.fixture
def db_session():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    db = factory()
    try:
        yield db
    finally:
        db.close()
        Base.metadata.drop_all(engine)


def add_scope(db: Session, *, config: dict | None = None, config_version: int = 1):
    user = User(username=f"owner-{id(db)}", password_hash="test")
    db.add(user)
    db.flush()
    workspace = Workspace(owner_id=user.id, name="Test")
    db.add(workspace)
    db.flush()
    knowledge_base = KnowledgeBase(
        workspace_id=workspace.id,
        name="KB",
        chunking_config=config or ChunkingConfig().model_dump(),
        chunking_config_version=config_version,
    )
    db.add(knowledge_base)
    db.flush()
    return workspace, knowledge_base


def add_document(db: Session, workspace: Workspace, knowledge_base: KnowledgeBase, filename: str) -> Document:
    document = Document(
        workspace_id=workspace.id,
        knowledge_base_id=knowledge_base.id,
        filename=filename,
        file_type=".md",
        storage_path=f"/{filename}",
        file_size=1,
        content_hash=f"hash-{filename}",
    )
    db.add(document)
    db.flush()
    return document


def test_docx_extraction_keeps_tables_in_document_order(tmp_path: Path):
    from docx import Document as DocxDocument

    source = DocxDocument()
    source.add_paragraph("Before table")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Policy"
    table.cell(0, 1).text = "Seven years"
    source.add_paragraph("After table")
    path = tmp_path / "policy.docx"
    source.save(path)

    content = extract_text(path, ".docx")

    assert content.index("Before table") < content.index("Policy | Seven years") < content.index("After table")


def test_ingestion_uses_selected_strategy_and_waits_for_every_document_reindex(db_session, monkeypatch):
    config = ChunkingConfig(
        strategy="heading",
        chunk_size=100,
        chunk_overlap=10,
        parent_chunk_size=512,
        enable_parent_child=True,
    )
    workspace, knowledge_base = add_scope(db_session, config=config.model_dump(), config_version=2)
    knowledge_base.chunking_needs_reindex = True
    first = add_document(db_session, workspace, knowledge_base, "first.md")
    second = add_document(db_session, workspace, knowledge_base, "second.md")
    old_chunk = DocumentChunk(
        workspace_id=workspace.id,
        knowledge_base_id=knowledge_base.id,
        document_id=first.id,
        sequence=0,
        content="Old indexed content",
        content_hash="old-hash",
        index_status="indexed",
    )
    db_session.add(old_chunk)
    db_session.commit()
    old_chunk_id = old_chunk.id

    text = (
        "# Alpha\n\n" + "alpha content with useful details. " * 30
        + "\n\n# Beta\n\n" + "beta content with separate details. " * 30
    )
    monkeypatch.setattr(ingestion, "extract_text", lambda *_args: text)

    stores = []

    class FakeVectorStore:
        def __init__(self):
            self.deleted: list[str] = []
            self.upserted: list[dict] = []
            stores.append(self)

        def delete(self, ids):
            self.deleted.extend(ids)

        def upsert(self, values):
            self.upserted.extend(values)

    async def fake_embed(_self, texts):
        return [[1.0] for _ in texts]

    monkeypatch.setattr(ingestion, "VectorStore", FakeVectorStore)
    monkeypatch.setattr(ingestion.EmbeddingProvider, "embed", fake_embed)

    ingestion.process_document(db_session, first.id)

    first_chunks = list(db_session.scalars(select(DocumentChunk).where(DocumentChunk.document_id == first.id)))
    parents = [chunk for chunk in first_chunks if chunk.chunk_level == "parent"]
    children = [chunk for chunk in first_chunks if chunk.chunk_level == "child"]
    assert parents and children
    assert all(chunk.chunking_strategy == "heading" for chunk in parents + children)
    assert all(chunk.parent_chunk_id in {parent.id for parent in parents} for chunk in children)
    assert {value["id"] for value in stores[0].upserted} == {chunk.id for chunk in children}
    assert old_chunk_id in stores[0].deleted
    assert knowledge_base.chunking_needs_reindex is True
    assert first.chunking_config_version == 2

    ingestion.process_document(db_session, second.id)

    assert knowledge_base.chunking_needs_reindex is False
    assert second.chunking_config_version == 2


def test_hybrid_search_retrieves_children_expands_parent_and_prunes_stale_vectors(db_session, monkeypatch):
    workspace, knowledge_base = add_scope(db_session)
    document = add_document(db_session, workspace, knowledge_base, "evidence.md")
    parent = DocumentChunk(
        workspace_id=workspace.id,
        knowledge_base_id=knowledge_base.id,
        document_id=document.id,
        sequence=0,
        content="Background context plus the recognizable evidence passage.",
        content_hash="parent-hash",
        chunk_level="parent",
        index_status="context",
    )
    db_session.add(parent)
    db_session.flush()
    child = DocumentChunk(
        workspace_id=workspace.id,
        knowledge_base_id=knowledge_base.id,
        document_id=document.id,
        sequence=0,
        content="The recognizable evidence passage.",
        content_hash="child-hash",
        chunk_level="child",
        parent_chunk_id=parent.id,
        index_status="indexed",
    )
    db_session.add(child)
    db_session.commit()

    class FakeVectorStore:
        deleted: list[str] = []

        def search(self, *_args):
            return [SimpleNamespace(id="stale-vector"), SimpleNamespace(id=child.id)]

        def delete(self, ids):
            self.deleted.extend(ids)

    class FakeEmbeddingProvider:
        async def embed(self, _texts):
            return [[1.0]]

    class FakeRerankProvider:
        async def rerank(self, _query, documents, _top_n):
            return [(index, 0.08) for index in range(len(documents))]

    store = FakeVectorStore()
    monkeypatch.setattr(retrieval, "EmbeddingProvider", FakeEmbeddingProvider)
    monkeypatch.setattr(retrieval, "RerankProvider", FakeRerankProvider)

    hits = asyncio.run(retrieval.hybrid_search(db_session, store, workspace.id, knowledge_base.id, "recognizable evidence"))

    assert len(hits) == 1
    assert hits[0]["chunk"].id == child.id
    assert hits[0]["context_chunk"].id == parent.id
    assert hits[0]["rank"] == 1
    assert hits[0]["rerank_score"] == pytest.approx(0.08)
    assert store.deleted == ["stale-vector"]
    context = retrieval.format_retrieval_context(hits)
    assert f"[source:{child.id}]" in context
    assert parent.content in context
